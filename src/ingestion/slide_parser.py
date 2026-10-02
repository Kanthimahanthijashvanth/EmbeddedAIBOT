from pptx import Presentation


class SlideParser:

    def __init__(self, file_path):
        self.file_path = file_path

    def parse(self):

        presentation = Presentation(self.file_path)

        slides = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            slide_data = {
                "slide_number": slide_number,
                "text": [],
                "images": [],
                "shapes": []
            }

            for shape in slide.shapes:

                # -------------------------
                # TEXT
                # -------------------------

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:

                        slide_data["text"].append({
                            "shape_id": shape.shape_id,
                            "name": shape.name,
                            "content": text
                        })

                # -------------------------
                # IMAGES
                # -------------------------

                if shape.shape_type == 13:

                    slide_data["images"].append({
                        "shape_id": shape.shape_id,
                        "name": shape.name
                    })

                # -------------------------
                # SHAPE INFORMATION
                # -------------------------

                slide_data["shapes"].append({
                    "shape_id": shape.shape_id,
                    "name": shape.name,
                    "type": str(shape.shape_type)
                })

            slides.append(slide_data)

        return slides