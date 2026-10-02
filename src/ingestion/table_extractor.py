from pptx import Presentation


class TableExtractor:

    def __init__(self, file_path):
        self.file_path = file_path

    def extract(self):

        presentation = Presentation(self.file_path)

        tables = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            for shape in slide.shapes:

                if not shape.has_table:
                    continue

                table = shape.table

                rows = []

                for row in table.rows:

                    cells = []

                    for cell in row.cells:
                        cells.append(
                            cell.text.strip()
                        )

                    rows.append(cells)

                tables.append({
                    "slide_number": slide_number,
                    "table_number": len(tables) + 1,
                    "rows": rows
                })

        return tables

    def to_markdown(self, tables):

        markdown_tables = []

        for table in tables:

            rows = table["rows"]

            if not rows:
                continue

            header = rows[0]

            markdown = []

            # Header
            markdown.append(
                "| " + " | ".join(header) + " |"
            )

            # Separator
            markdown.append(
                "| "
                + " | ".join(
                    ["---"] * len(header)
                )
                + " |"
            )

            # Data rows
            for row in rows[1:]:

                # Make sure row has same number of cells
                row = row + [""] * (
                    len(header) - len(row)
                )

                markdown.append(
                    "| "
                    + " | ".join(row[:len(header)])
                    + " |"
                )

            markdown_tables.append({
                "slide_number": table["slide_number"],
                "table_number": table["table_number"],
                "content": "\n".join(markdown)
            })

        return markdown_tables