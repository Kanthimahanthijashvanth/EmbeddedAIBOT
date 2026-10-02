from pathlib import Path

from langchain_core.documents import Document
from pptx import Presentation
from pypdf import PdfReader


class WorkshopLoader:

    SUPPORTED_EXTENSIONS = {".pptx", ".pdf"}

    def __init__(self, file_path):
        self.file_path = Path(file_path)

    def validate(self):
        if not self.file_path.exists():
            raise FileNotFoundError(
                f"File not found: {self.file_path}"
            )

        if self.file_path.suffix.lower() not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type: {self.file_path.suffix}"
            )

    def load(self):
        self.validate()

        extension = self.file_path.suffix.lower()

        if extension == ".pptx":
            return self._load_pptx()

        if extension == ".pdf":
            return self._load_pdf()

    def _load_pptx(self):

        presentation = Presentation(str(self.file_path))

        documents = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            slide_text = []

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:
                        slide_text.append(text)

            content = "\n".join(slide_text)

            documents.append(
                Document(
                    page_content=content,
                    metadata={
                        "source": str(self.file_path),
                        "file_name": self.file_path.name,
                        "page_number": slide_number,
                        "content_type": "pptx",
                        "workshop_id": self.file_path.stem
                    }
                )
            )

        return documents

    def _load_pdf(self):

        reader = PdfReader(str(self.file_path))

        documents = []

        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):

            content = page.extract_text() or ""

            documents.append(
                Document(
                    page_content=content,
                    metadata={
                        "source": str(self.file_path),
                        "file_name": self.file_path.name,
                        "page_number": page_number,
                        "content_type": "pdf",
                        "workshop_id": self.file_path.stem
                    }
                )
            )

        return documents