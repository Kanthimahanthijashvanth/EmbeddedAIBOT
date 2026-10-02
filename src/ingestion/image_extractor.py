from pathlib import Path
from pptx import Presentation


class ImageExtractor:

    def __init__(self, file_path, output_dir="output/extracted_images"):
        self.file_path = Path(file_path)
        self.output_dir = Path(output_dir)

        self.output_dir.mkdir(
            parents=True,
            exist_ok=True
        )

    def extract(self):

        presentation = Presentation(str(self.file_path))

        extracted_images = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            slide_images = []

            for image_number, shape in enumerate(
                slide.shapes,
                start=1
            ):

                if shape.shape_type == 13:  # Picture

                    image = shape.image

                    extension = image.ext

                    filename = (
                        f"{self.file_path.stem}"
                        f"_slide_{slide_number}"
                        f"_image_{image_number}"
                        f".{extension}"
                    )

                    output_path = self.output_dir / filename

                    with open(output_path, "wb") as f:
                        f.write(image.blob)

                    image_info = {
                        "slide_number": slide_number,
                        "image_number": image_number,
                        "file_name": filename,
                        "file_path": str(output_path),
                        "extension": extension
                    }

                    slide_images.append(image_info)
                    extracted_images.append(image_info)

        return extracted_images